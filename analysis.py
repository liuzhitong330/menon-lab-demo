#!/usr/bin/env python3
"""Reproduce a readout-only audit of eight published, directly paired sheets.

No participant identifiers are inferred, no sheets are joined, and no raw MRI,
ABCD, or MAT data are used. The seventh and eighth comparisons are not treated
as independent participant samples. Run with Python 3.12 and requirements.txt.
"""

from __future__ import annotations

import argparse
import contextlib
import hashlib
import io
import json
import math
from pathlib import Path
from statistics import fmean
from urllib.request import Request, urlopen
from unittest.mock import patch

import openpyxl

ROOT = Path(__file__).resolve().parent
WORKBOOK_NAME = "41467_2024_46511_MOESM3_ESM.xlsx"
WORKBOOK_SHA256 = "469d69d4303c7321c093f19d5f390293fb55b10cb6b02673ca7afc34269f90ee"
WORKBOOK_URL = "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC10928152/supplementaryFiles"
PAPER_URL = "https://www.nature.com/articles/s41467-024-46511-5"
UPSTREAM_REPO = "https://github.com/scsnl/Cai_Multiple_Demand_System_2023"
UPSTREAM_COMMIT = "cfe1931b2ac317c992642a0011eedac33cad2534"
UPSTREAM_PATH = "brainbehav/brainbehav_cca.py"
UPSTREAM_SHA256 = "d880fc28f80fe4c4b69508e36f7ac1bcfc871f33ae49678916e8be474de9df1b"
UPSTREAM_URL = f"https://raw.githubusercontent.com/scsnl/Cai_Multiple_Demand_System_2023/{UPSTREAM_COMMIT}/{UPSTREAM_PATH}"
ALIGNMENT_TOLERANCE = 5e-10

# Frozen before the audit. N means finite direct A/B pairs in that exact sheet,
# not the nominal cohort N in a paper caption, and not unique people pooled.
DATASETS = (
    ("axcpt", "Figure3e", "AX-CPT", "DMCC", 50),
    ("cuedts", "Figure3f", "Cued task switching", "DMCC", 49),
    ("sternberg", "Figure3g", "Sternberg", "DMCC", 50),
    ("stroop", "Figure3h", "Stroop", "DMCC", 48),
    ("sst", "Figure5b", "Stop-signal task", "Stanford SST", 44),
    ("rp1", "Figure7c", "Relational processing: session 1", "HCP", 90),
    ("rp2", "Figure7d", "Relational processing: session 2", "HCP", 87),
    ("sst_neurosynth", "FigureS22b", "Stop-signal task: NeuroSynth ROIs", "Stanford SST", 44),
)


def checked_bytes(path: Path, expected_sha256: str) -> bytes:
    data = Path(path).read_bytes()
    actual = hashlib.sha256(data).hexdigest()
    if actual != expected_sha256:
        raise ValueError(f"Checksum mismatch for {Path(path).name}: {actual}")
    return data


def checked_download(url: str, expected_sha256: str) -> bytes:
    """Only callers' fixed, reviewed URLs are used; no arbitrary-URL CLI."""
    with urlopen(Request(url, headers={"User-Agent": "MenonReadoutAudit/1.0"}), timeout=30) as response:
        data = response.read(2_000_000)
    actual = hashlib.sha256(data).hexdigest()
    if actual != expected_sha256:
        raise ValueError(f"Downloaded source checksum mismatch: {actual}")
    return data


def number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def extract_pairs(workbook_path: Path) -> dict[str, list[dict]]:
    checked_bytes(workbook_path, WORKBOOK_SHA256)
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    extracted = {}
    try:
        for _, sheet_name, _, _, expected_n in DATASETS:
            sheet = workbook[sheet_name]
            if sheet.max_column != 2 or tuple(sheet.values)[0] != ("brain", "behavior"):
                raise ValueError(f"Unexpected columns in {sheet_name}")
            pairs = []
            for source_row, (x, y) in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                if x is None and y is None:
                    continue
                # Fail loudly on a changed or incomplete pair instead of dropping
                # observations silently. All frozen source sheets are complete.
                if not number(x) or not number(y):
                    raise ValueError(f"Invalid pair at {sheet_name} Excel row {source_row}")
                if not 0 <= x <= 1:
                    raise ValueError(f"Occupancy is not a fraction at {sheet_name}:{source_row}")
                pairs.append({"source_row": source_row, "x": x, "y": y})
            if len(pairs) != expected_n:
                raise ValueError(f"Expected {expected_n} pairs in {sheet_name}, got {len(pairs)}")
            extracted[sheet_name] = pairs
    finally:
        workbook.close()
    return extracted


def fit(x: list[float], y: list[float]) -> tuple[float, float, float]:
    """Centered OLS with intercept and Pearson r; no significance calculation."""
    if len(x) != len(y) or len(x) < 3:
        raise ValueError("Need at least three paired observations")
    if not all(number(value) for value in x + y):
        raise ValueError("All observations must be finite numbers")
    mx, my = fmean(x), fmean(y)
    xx = math.fsum((value - mx) ** 2 for value in x)
    yy = math.fsum((value - my) ** 2 for value in y)
    xy = math.fsum((a - mx) * (b - my) for a, b in zip(x, y))
    if xx <= 0 or yy <= 0:
        raise ValueError("Constant input or outcome is not analyzable")
    slope = xy / xx
    return slope, my - slope * mx, xy / math.sqrt(xx * yy)


def error_metrics(y: list[float], predicted: list[float]) -> tuple[float, float]:
    if len(y) != len(predicted) or not y:
        raise ValueError("Error vectors must have equal nonzero lengths")
    errors = [actual - prediction for actual, prediction in zip(y, predicted)]
    return math.sqrt(fmean(value * value for value in errors)), fmean(abs(value) for value in errors)


def calculated(value: float) -> float:
    """Stabilize derived output; original source x/y values are never rounded."""
    return round(float(value), 12)


def analyze_pairs(pairs: list[dict]) -> tuple[list[dict], dict]:
    x, y = [row["x"] for row in pairs], [row["y"] for row in pairs]
    slope, intercept, r = fit(x, y)
    rows, predictions, baselines, omitted_correlations = [], [], [], []
    for index, source in enumerate(pairs):
        train_x, train_y = x[:index] + x[index + 1:], y[:index] + y[index + 1:]
        train_slope, train_intercept, train_r = fit(train_x, train_y)
        prediction = train_intercept + train_slope * x[index]
        baseline = fmean(train_y)  # Never use the held-out outcome in this mean.
        predictions.append(prediction)
        baselines.append(baseline)
        omitted_correlations.append(train_r)
        rows.append({
            **source,
            "yhat": calculated(intercept + slope * x[index]),
            "loo_pred": calculated(prediction),
            "baseline_pred": calculated(baseline),
            "loo_r": calculated(train_r),
        })
    rmse, mae = error_metrics(y, predictions)
    baseline_rmse, baseline_mae = error_metrics(y, baselines)
    metrics = {
        "r": r, "slope": slope, "intercept": intercept,
        "rmse": rmse, "baseline_rmse": baseline_rmse,
        "rmse_improvement": 100 * (1 - rmse / baseline_rmse),
        "mae": mae, "baseline_mae": baseline_mae,
        "mae_improvement": 100 * (1 - mae / baseline_mae),
        "loo_r_min": min(omitted_correlations),
        "loo_r_max": max(omitted_correlations),
    }
    return rows, {key: calculated(value) for key, value in metrics.items()}


def check_sst_alignment(primary: list[dict], alternative: list[dict]) -> dict:
    """Numerical same-order check only: this does not recover participant IDs."""
    if len(primary) != 44 or len(alternative) != 44:
        raise ValueError("SST sensitivity requires the frozen 44 pairs in both sheets")
    a, b = [row["y"] for row in primary], [row["y"] for row in alternative]
    unique_primary = len(set(a))
    unique_alternative = len(set(b))
    if unique_primary != 44 or unique_alternative != 44:
        raise ValueError("Outcome values must be unique within both SST sheets")
    differences = [abs(x - y) for x, y in zip(a, b)]
    if max(differences) > ALIGNMENT_TOLERANCE:
        raise ValueError("SST outcomes no longer match in source order")
    return {
        "primary_sheet": "Figure5b", "alternative_sheet": "FigureS22b",
        "n": 44, "same_order_matches": sum(d <= ALIGNMENT_TOLERANCE for d in differences),
        "unique_primary_values": unique_primary,
        "unique_alternative_values": unique_alternative,
        "absolute_tolerance": ALIGNMENT_TOLERANCE,
        "maximum_absolute_difference": max(differences),
        "interpretation": "44 unique outcome values match in original source order within tolerance; no participant IDs are inferred. Each sheet retains its own literal outcome values.",
    }


def run_upstream_cca(pairs: list[dict], source_path: Path | None = None) -> dict:
    """Execute hash-verified upstream code on one published one-column pair.

    The unlicensed upstream source is fetched at a fixed commit, not bundled.
    Its four configuration assignments are the ONLY source substitutions.
    Its CSV reader receives in-memory frames; no raw/MAT data or joins occur.
    This reviewed script only imports libraries, fits CCA, and prints a result.
    The source hash is checked before execution, including for a local path.
    """
    import matplotlib
    import pandas as pd

    matplotlib.use("Agg")
    source = (checked_bytes(source_path, UPSTREAM_SHA256) if source_path
              else checked_download(UPSTREAM_URL, UPSTREAM_SHA256)).decode("utf-8")
    replacements = {
        "state_OR_fname = ''": "state_OR_fname = 'in_memory_brain.csv'",
        "behav_fname = ''": "behav_fname = 'in_memory_behavior.csv'",
        "X_varnames = []": "X_varnames = ['brain']",
        "y_varnames = []": "y_varnames = ['behavior']",
    }
    for original, replacement in replacements.items():
        if source.count(original) != 1:
            raise ValueError(f"Unexpected upstream configuration: {original}")
        source = source.replace(original, replacement, 1)
    frames = {
        "in_memory_brain.csv": pd.DataFrame({"brain": [row["x"] for row in pairs]}),
        "in_memory_behavior.csv": pd.DataFrame({"behavior": [row["y"] for row in pairs]}),
    }

    def in_memory_csv(filename, *args, **kwargs):
        if filename not in frames or args or kwargs:
            raise ValueError("Unexpected upstream input request")
        return frames[filename].copy()

    namespace = {"__name__": "__reviewed_pinned_cca__"}
    with patch.object(pd, "read_csv", side_effect=in_memory_csv), contextlib.redirect_stdout(io.StringIO()):
        exec(compile(source, UPSTREAM_URL, "exec"), namespace)
    # Upstream prints a p-value as part of scipy's result; it is deliberately
    # not saved or presented. Only the first canonical correlation is retained.
    canonical_r = float(namespace["spss"].pearsonr(namespace["X_c"], namespace["y_c"]).statistic)
    _, _, pearson_r = fit([row["x"] for row in pairs], [row["y"] for row in pairs])
    if not math.isclose(abs(canonical_r), abs(pearson_r), rel_tol=0, abs_tol=1e-10):
        raise ValueError("One-column upstream CCA does not agree with absolute Pearson r")
    return {
        "repository": UPSTREAM_REPO, "commit": UPSTREAM_COMMIT,
        "path": UPSTREAM_PATH, "source_sha256": UPSTREAM_SHA256,
        "source_url": UPSTREAM_URL, "input_sheet": "Figure3e", "n": len(pairs),
        "execution": "Hash-verified upstream source executed in memory; only its two filename and two variable-list assignments were substituted.",
        "first_canonical_correlation": calculated(canonical_r),
        "absolute_pearson_r": calculated(abs(pearson_r)),
        "absolute_difference": abs(abs(canonical_r) - abs(pearson_r)),
        "scope": "One-column CCA recipe check only; not a replication of the original multivariate CCA, BSDS, feature selection, or full imaging pipeline.",
        "license": "No upstream repository license found at the pinned commit; upstream source is not redistributed in this repository.",
    }


def build_analysis(workbook_path: Path, upstream_source: Path | None = None) -> dict:
    extracted = extract_pairs(workbook_path)
    datasets = []
    for dataset_id, sheet, label, cohort, _ in DATASETS:
        rows, metrics = analyze_pairs(extracted[sheet])
        datasets.append({
            "id": dataset_id, "sheet": sheet, "label": label, "cohort": cohort,
            "n": len(rows), "xlabel": "Shared-state occupancy (fraction)",
            "ylabel": "Published performance index (source units)",
            "rows": rows, "metrics": metrics,
        })
    return {
        "datasets": datasets,
        "provenance": {
            "schema_version": 1,
            "analysis_plan_date": "2026-10-30",
            "publication": "Cai, Taghia & Menon (2024), Nature Communications 15, 2185",
            "doi": "10.1038/s41467-024-46511-5", "paper_url": PAPER_URL,
            "source_workbook": WORKBOOK_NAME, "source_url": WORKBOOK_URL,
            "source_delivery": "The Europe PMC supplementaryFiles URL returns an archive. The named Source Data workbook was extracted unchanged and is bundled under inputs/.",
            "source_sha256": WORKBOOK_SHA256,
            "source_license": "CC BY 4.0", "license_url": "https://creativecommons.org/licenses/by/4.0/",
            "attribution": "Source data: Cai, Taghia & Menon (2024). Original workbook unchanged. Derived readout audit and interactive presentation are new, assistant-supported analyses.",
            "analysis_type": "Published processed-pair readout audit; no raw-data reanalysis.",
            "primary_sheets": [row[1] for row in DATASETS[:-1]],
            "prespecified_sensitivity_sheet": "FigureS22b",
            "input_columns": {"A": "brain: shared-state occupancy fraction", "B": "behavior: published task-specific performance index; higher means better performance"},
            "outcome_definitions": {
                "source": "Supplementary Table S9, page 44; source-sheet numerical values retained without renormalization.",
                "Figure3e": "AX-CPT: (AY RT - BX RT) / (AY RT + BX RT)",
                "Figure3f": "Cued task switching: congruent RT - incongruent RT",
                "Figure3g": "Sternberg: high-load accuracy / high-load RT",
                "Figure3h": "Stroop: congruent RT - incongruent RT",
                "Figure5b": "SST: 1 / stop-signal reaction time (SSRT)",
                "Figure7c": "Relational processing session 1: relational accuracy / relational RT",
                "Figure7d": "Relational processing session 2: relational accuracy / relational RT",
                "FigureS22b": "SST: 1 / stop-signal reaction time (SSRT), alternative NeuroSynth ROI definition",
            },
            "row_identity": "source_row is the one-based Excel row in the named sheet, not a participant ID. No cross-task or cross-session matching is inferred.",
            "inclusion": "All finite directly paired A/B values already exported in the eight frozen sheets are retained by this audit. The authors report excluding points at least 3 SD from the mean in their univariate analyses (Supplementary Methods, section H, page 8); the exact exclusion mask is not reconstructed. Counts are sheet-specific, not pooled participant N.",
            "method": "Pearson r and OLS with intercept. For each source row, refit the linear readout on all other rows and predict that row; compare to the other rows' mean outcome. RMSE and MAE use all held-out predictions.",
            "improvement_definition": "100 × (1 − readout error / training-only mean baseline error); positive means lower error, negative means higher error. Errors remain in each source outcome's units.",
            "single_row_diagnostic": "loo_r is the correlation after omitting that source row. Its minimum–maximum is a jackknife sensitivity range, not a confidence interval. No omission changes the reported full-data results.",
            "numeric_policy": "Source x/y numeric values are retained literally. Derived row predictions and summary metrics are rounded to 12 decimal places for deterministic display.",
            "sst_alignment": check_sst_alignment(extracted["Figure5b"], extracted["FigureS22b"]),
            "upstream_check": run_upstream_cca(extracted["Figure3e"], upstream_source),
            "limitations": [
                "The authors report a threshold of at least 3 SD from the mean for point exclusion in univariate analyses. This audit retains all already exported pairs, does not reconstruct the exact exclusion mask, and does not refit that original filtering within held-out folds.",
                "Published states/features and behavioral indices were selected in the original study; the state model may include all participants. Only the linear readout is held out here.",
                "The outcome is a published task-specific behavioral index, not a CCA-derived composite. The supplied normalization is not fully documented for every task; values remain in source units.",
                "Overlapping tasks, sessions, and the alternate ROI analysis are not independent replications; row IDs cannot establish participant overlap.",
                "A one-row omission range measures sensitivity, not sampling uncertainty or a confidence interval.",
                "These between-person published pairs do not test within-person mechanisms, clinical utility, causality, or external generalization.",
                "Cross-task outcome definitions differ; do not pool rows or rank tasks by absolute error units.",
            ],
        },
    }


def serialize(payload: dict) -> str:
    return json.dumps(payload, ensure_ascii=False, allow_nan=False, indent=2) + "\n"


def javascript(payload: dict) -> str:
    return "// Generated by analysis.py. Do not edit by hand.\nwindow.MENON_DATA = " + serialize(payload).rstrip() + ";\n"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=ROOT / "inputs" / WORKBOOK_NAME)
    parser.add_argument("--upstream-source", type=Path, help="Optional local upstream script; its fixed SHA-256 is still required")
    parser.add_argument("--output-dir", type=Path, default=ROOT)
    parser.add_argument("--check", action="store_true", help="Recompute and verify checked-in outputs without writing")
    args = parser.parse_args()
    payload = build_analysis(args.workbook, args.upstream_source)
    outputs = {"analysis.json": serialize(payload), "data.js": javascript(payload)}
    for filename, content in outputs.items():
        path = args.output_dir / filename
        if args.check:
            if path.read_text(encoding="utf-8") != content:
                raise SystemExit(f"Generated output differs: {path.name}")
        else:
            args.output_dir.mkdir(parents=True, exist_ok=True)
            path.write_text(content, encoding="utf-8")
    print("Verified" if args.check else "Generated", "8 paired-sheet audits and one pinned upstream CCA check.")
    for dataset in payload["datasets"]:
        metric = dataset["metrics"]
        print(f"{dataset['sheet']:12} n={dataset['n']:2} r={metric['r']:.6f} RMSE change={metric['rmse_improvement']:+.3f}% MAE change={metric['mae_improvement']:+.3f}%")


if __name__ == "__main__":
    main()
